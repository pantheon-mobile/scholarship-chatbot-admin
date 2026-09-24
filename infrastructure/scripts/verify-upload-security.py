"""Bounded, development-only F-003/F-004 regression checks. No credentials logged."""
import io
import json
import subprocess
import time
import uuid
from pathlib import Path

import requests
from openpyxl import load_workbook

HOST = 'https://scholarship-chatbot-dev.pantheon-mobile.com'
STACK = 'ScholarshipChatbot-development'
MIB = 1024 * 1024
results = []

def aws(*args):
    return subprocess.check_output(['aws', *args, '--region', 'ap-northeast-1', '--output', 'json'], text=True)


def check(label, condition):
    results.append((label, bool(condition)))
    print(('PASS ' if condition else 'FAIL ') + label, flush=True)
    if not condition:
        raise AssertionError(label)


def workbook_bytes(book):
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def main():
    outputs = json.loads(aws('cloudformation', 'describe-stacks', '--stack-name', STACK))['Stacks'][0]['Outputs']
    outputs = {x['OutputKey']: x['OutputValue'] for x in outputs}
    check('Development host guard', outputs['ApplicationUrl'].rstrip('/') == HOST)
    secret = json.loads(aws('secretsmanager', 'get-secret-value', '--secret-id', outputs['CpfLoginPasswordSecretName']))['SecretString']
    session = requests.Session()
    def request(method, path, **kwargs):
        response = session.request(method, HOST + '/api/v1' + path, timeout=90, **kwargs)
        response.raise_for_status()
        return response
    token = request('POST', '/auth/development/token', json={
        'subject': 'upload-security-verification', 'display_name': '受信制限検証', 'role': 'admin', 'password': secret,
    }).json()['token']
    del secret
    request('POST', '/auth/development/cpf', json={'token': token})
    del token
    check('Development login/session', request('GET', '/auth/session').status_code == 200)
    marker = 'security-check-' + uuid.uuid4().hex
    try:
        # Never expose test content as an answer source.
        row = request('POST', '/data-sources/files', files={
            'files': (marker + '.txt', 'アップロード検証用です。'.encode(), 'text/plain'),
        }, data={'answer_source_enabled': 'false', 'reference_link_visible': 'false'}).json()['items'][0]
        check('Normal file upload', row['status'] == 'PREPARING' and not row['answer_source_enabled'])
        book = load_workbook(io.BytesIO(request('GET', '/data-sources/export', params={'keyword': marker}).content))
        check('Export contains only test source', book.active.max_row == 2 and book.active.cell(2, 1).value == row['id'])
        imported = request('POST', '/data-sources/import', files={'file': ('sources.xlsx', workbook_bytes(book))}).json()
        check('Data source Excel import', imported['processed_count'] == 1)
        book = load_workbook(io.BytesIO(request('GET', '/faqs/import-template').content))
        values = [None] * book.active.max_column
        values[1], values[2], values[-1] = marker, '検証用FAQです。', '非公開'
        book.active.append(values)
        imported = request('POST', '/faqs/import', files={'file': ('faq.xlsx', workbook_bytes(book))}).json()
        check('FAQ Excel import', imported['created_count'] == 1)
        answer = request('POST', '/chat/messages', json={'question': '登録資料に情報がない場合の動作確認です'}).json()
        check('Chat response', isinstance(answer.get('answer'), str) and bool(answer['answer']))
        # Approximately 11 MiB is sent sequentially; no resource exhaustion test.
        def oversized():
            yield b'--security-probe\r\nContent-Disposition: form-data; name="file"; filename="probe.xlsx"\r\n\r\n'
            for _ in range(11):
                yield b'x' * MIB
            yield b'x\r\n--security-probe--\r\n'
        response = requests.post(HOST + '/api/v1/faqs/import', data=oversized(),
            headers={'Content-Type': 'multipart/form-data; boundary=security-probe'}, timeout=45)
        check('Unauthenticated chunked body limit / Japanese error', response.status_code == 413 and
              response.json()['detail']['code'] == 'REQUEST_BODY_TOO_LARGE' and
              '上限' in response.json()['detail']['message'])
        for media in ['multipart', 'urlencoded']:
            if media == 'multipart':
                response = requests.post(HOST + '/api/v1/faqs/import', files={'field': (None, 'a' * (MIB + 1))}, timeout=45)
            else:
                response = requests.post(HOST + '/api/v1/faqs/import', data={'field': 'a' * (MIB + 1)}, timeout=45)
            check('Patched parser rejects oversized ' + media + ' field', response.status_code == 400)
        check('Health after bounded rejection checks', request('GET', '/health').json()['status'] == 'ok')
    finally:
        # Locate only this run's uniquely named FAQ, including when an import response was lost.
        try:
            faqs = request('GET', '/faqs', params={'keyword': marker}).json()['items']
            for faq in faqs:
                if faq['question'] == marker:
                    request('DELETE', '/faqs/' + str(faq['id']), params={'version': faq['version']})
            check('Test FAQ cleanup', not request('GET', '/faqs', params={'keyword': marker}).json()['items'])
        finally:
            sources = request('GET', '/data-sources', params={'keyword': marker}).json()['items']
            for source in sources:
                if source.get('file', {}).get('file_name') != marker + '.txt':
                    continue
                # A scheduled worker may have claimed the source meanwhile.
                for _ in range(60):
                    source = request('GET', '/data-sources/' + str(source['id'])).json()
                    if source['status'] != 'TRAINING':
                        break
                    time.sleep(5)
                request('DELETE', '/data-sources/' + str(source['id']), params={'version': source['version']})
            check('Test source cleanup', not request('GET', '/data-sources', params={'keyword': marker}).json()['items'])


if __name__ == '__main__':
    try:
        main()
    finally:
        Path('upload-security-report.md').write_text('# Development upload security verification\n\n' +
            '\n'.join(f'- {"PASS" if ok else "FAIL"}: {label}' for label, ok in results) + '\n')
