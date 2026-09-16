from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from app.services.excel_format import apply_download_format

# Customer form-definition rows 6-86; widths use column E, not pixel conversion.
EXPECTED = {'データソース一覧': [['ID', 9.3, False, 'right'], ['種類', 9.3, False, 'left'], ['タイトル', 59.3, True, 'left'], ['ファイル名／URL', 59.3, True, 'left'], ['形式', 7.3, False, 'left'], ['状態', 7.3, False, 'left'], ['カテゴリ', 24.3, True, 'left'], ['タイプ', 14.3, True, 'left'], ['年度', 14.3, True, 'left'], ['対象', 14.3, True, 'left'], ['サイズ', 9.3, False, 'right'], ['文字数', 9.3, False, 'right'], ['回答ソース', 12.3, False, 'left'], ['優先度', 12.3, False, 'left'], ['参照元リンク', 12.3, False, 'left'], ['更新日時', 19.3, False, 'right']], 'URLリスト': [['URL', 59.3, False, 'left'], ['タイトル', 59.3, True, 'left']], '種別': [['種別', 7.3, False, 'left'], ['種別ラベル名', 14.3, True, 'left'], ['種別値', 49.3, True, 'left']], 'FAQ一覧': [['ID', 9.3, False, 'right'], ['質問', 59.3, True, 'left'], ['回答', 79.3, True, 'left'], ['類似質問1', 39.3, True, 'left'], ['類似質問2', 39.3, True, 'left'], ['類似質問3', 39.3, True, 'left'], ['類似質問4', 39.3, True, 'left'], ['類似質問5', 39.3, True, 'left'], ['類似質問6', 39.3, True, 'left'], ['類似質問7', 39.3, True, 'left'], ['類似質問8', 39.3, True, 'left'], ['類似質問9', 39.3, True, 'left'], ['類似質問10', 39.3, True, 'left'], ['問合せ区分', 24.3, True, 'left'], ['質問区分', 24.3, True, 'left'], ['年度', 24.3, True, 'left'], ['対象', 24.3, True, 'left'], ['チャット利用', 14.3, False, 'left'], ['更新日時', 19.3, False, 'right']], 'FAQ一括登録更新': [['ID', 9.3, False, 'right'], ['質問', 59.3, True, 'left'], ['回答', 79.3, True, 'left'], ['類似質問1', 39.3, True, 'left'], ['類似質問2', 39.3, True, 'left'], ['類似質問3', 39.3, True, 'left'], ['類似質問4', 39.3, True, 'left'], ['類似質問5', 39.3, True, 'left'], ['類似質問6', 39.3, True, 'left'], ['類似質問7', 39.3, True, 'left'], ['類似質問8', 39.3, True, 'left'], ['類似質問9', 39.3, True, 'left'], ['類似質問10', 39.3, True, 'left'], ['問合せ区分', 24.3, True, 'left'], ['質問区分', 24.3, True, 'left'], ['年度', 24.3, True, 'left'], ['対象', 24.3, True, 'left'], ['チャット利用', 14.3, False, 'left']], '区分': [['区分', 7.3, False, 'right'], ['区分ラベル名', 14.3, True, 'left'], ['区分値', 49.3, True, 'left']], 'カテゴリ一覧': [['ID', 9.3, False, 'right'], ['カテゴリ1', 49.3, True, 'left'], ['カテゴリ2', 49.3, True, 'left'], ['カテゴリ3', 49.3, True, 'left']], 'チャット履歴': [['チャットID', 24.3, True, 'left'], ['ログインID', 19.3, True, 'left'], ['権限', 19.3, True, 'left'], ['応答ID', 24.3, True, 'left'], ['応答No', 7.3, False, 'right'], ['質問', 59.3, True, 'left'], ['回答', 79.3, True, 'left'], ['回答種別', 9.3, False, 'left'], ['評価', 9.3, False, 'left'], ['コメント', 39.3, True, 'left'], ['質問受付日時', 19.3, False, 'right'], ['回答完了日時', 19.3, False, 'right']], 'ユーザリスト': [['ログインID', 19.3, True, 'left'], ['権限', 19.3, True, 'left'], ['氏名', 29.3, True, 'left'], ['最終アクセス日時', 19.3, False, 'right']]}


@pytest.mark.parametrize("name,columns", EXPECTED.items())
def test_download_format_survives_excel_serialization(name, columns):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = name
    sheet.append([column[0] for column in columns])
    values = ["文字列\n改行" for _ in columns]
    sheet.append(values)
    apply_download_format(sheet)
    output = BytesIO()
    workbook.save(output)
    saved = load_workbook(BytesIO(output.getvalue())).active
    assert list(saved.values)[1] == tuple(values)
    for index, (_, width, wrap, horizontal) in enumerate(columns, start=1):
        dimension = saved.column_dimensions[get_column_letter(index)]
        assert dimension.width == width
        assert dimension.font.name == "游ゴシック"
        assert bool(dimension.alignment.wrap_text) == wrap
        header, data = saved.cell(1, index), saved.cell(2, index)
        assert header.fill.fgColor.rgb[-6:] == "F8F7F7"
        assert header.alignment.horizontal == header.alignment.vertical == "center"
        assert data.alignment.horizontal == horizontal
        assert data.alignment.vertical == "top"
        assert bool(data.alignment.wrap_text) == wrap
        for cell in (header, data):
            assert cell.font.name == "游ゴシック"
            assert cell.font.sz == 11 and not cell.font.bold
            assert cell.font.color.rgb[-6:] == "000000"
            assert all(getattr(cell.border, side) is None or getattr(cell.border, side).style is None for side in ("left", "right", "top", "bottom"))


def test_row_heights_handle_japanese_wrapping_newlines_and_excel_limit():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "URLリスト"
    sheet.append(["URL", "タイトル"])
    sheet.append(["https://example.com", "短いタイトル"])
    sheet.append(["https://example.com", "一行目\n二行目\n三行目"])
    sheet.append(["https://example.com", "長い日本語のタイトル" * 30])
    sheet.append(["https://example.com", "長文" * 2000])
    apply_download_format(sheet)
    output = BytesIO()
    workbook.save(output)
    saved = load_workbook(BytesIO(output.getvalue())).active
    assert saved.sheet_format.defaultRowHeight == 18
    assert saved.row_dimensions[1].height == 18
    assert saved.row_dimensions[2].height == 18
    assert saved.row_dimensions[3].height == 54
    assert 54 < saved.row_dimensions[4].height <= 409.5
    assert saved.row_dimensions[5].height == 409.5
    assert saved.cell(5, 2).value == "長文" * 2000
