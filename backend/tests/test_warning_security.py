from io import BytesIO
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook, load_workbook
from fastapi import FastAPI
from app.middleware.csrf import CsrfMiddleware
from app.services.excel_format import append_safe_row, safe_csv_value
from app.services.client_ip import client_ip
from app.services.public_http import PublicHTTPSConnection, public_addresses
from app.services.citation_validation import safe_citation_uri


def test_export_strings_cannot_become_formulas_and_controls_are_removed():
    book = Workbook()
    append_safe_row(book.active, ['=1+1', '=HYPERLINK("https://example.com")', 'text\x01\t\n', 42])
    output = BytesIO()
    book.save(output)
    row = next(load_workbook(BytesIO(output.getvalue())).active.rows)
    assert [c.data_type for c in row] == ['s', 's', 's', 'n']
    assert row[0].value == '=1+1'
    assert row[2].value == 'text\t\n'
    assert safe_csv_value(' \t=1+1') == "' \t=1+1"
    assert safe_csv_value(42) == 42


@pytest.mark.parametrize('uri', ['javascript:alert(1)', 'data:text/html,x', '//evil.test/', 'https://user:pass@example.com', 'https://example.com\\@evil.test'])
def test_unsafe_citations(uri):
    assert safe_citation_uri(uri) is None


def test_safe_citations():
    for uri in ['https://example.com/doc', '/api/v1/chat/sources/1/download']:
        assert safe_citation_uri(uri) == uri


def test_forwarded_ip_only_from_trusted_proxy(monkeypatch):
    monkeypatch.setenv('TRUSTED_PROXY_CIDRS', '10.0.0.0/8')
    request = SimpleNamespace(client=SimpleNamespace(host='10.0.0.2'), headers={'x-forwarded-for': '1.2.3.4, 8.8.8.8'})
    assert client_ip(request) == '8.8.8.8'
    request.client.host = '9.9.9.9'
    assert client_ip(request) == '9.9.9.9'


def test_ssrf_blocks_mixed_public_private_resolution(monkeypatch):
    monkeypatch.setattr('socket.getaddrinfo', lambda *a, **k: [(0, 0, 0, '', ('8.8.8.8', 443)), (0, 0, 0, '', ('127.0.0.1', 443))])
    with pytest.raises(ValueError):
        public_addresses('example.com', 443)


def test_ssrf_pins_socket_ip_preserving_tls_hostname(monkeypatch):
    monkeypatch.setattr('app.services.public_http.public_addresses', lambda *a: ['8.8.8.8'])
    connect = Mock()
    monkeypatch.setattr('app.services.public_http.create_connection', connect)
    connection = PublicHTTPSConnection('example.com', 443)
    connection._new_conn()
    assert connect.call_args.args[0] == ('8.8.8.8', 443)
    assert connection.host == 'example.com'


@pytest.mark.anyio
async def test_csrf_rejects_foreign_and_missing_origin_with_cookie(monkeypatch):
    monkeypatch.setenv('CORS_ALLOWED_ORIGINS', 'https://chat.example.com')
    app = FastAPI()
    app.add_middleware(CsrfMiddleware)
    @app.post('/write')
    def write(): return {'ok': True}
    async with AsyncClient(transport=ASGITransport(app=app), base_url='https://chat.example.com', cookies={'scholarship_session': 'test'}) as client:
        for headers in [{}, {'Origin': 'https://evil.example.com'}, {'Origin': 'null'}]:
            assert (await client.post('/write', headers=headers)).status_code == 403
        assert (await client.post('/write', headers={'Origin': 'https://chat.example.com'})).status_code == 200


@pytest.mark.parametrize('address', ['127.0.0.1', '10.0.0.1', '169.254.169.254', '169.254.170.2', '224.0.0.1', '::1', 'ff02::1', '::ffff:127.0.0.1'])
def test_ssrf_special_addresses_rejected(address, monkeypatch):
    monkeypatch.setattr('socket.getaddrinfo', lambda *a, **k: [(0, 0, 0, '', (address, 443))])
    with pytest.raises(ValueError):
        public_addresses('example.com', 443)


@pytest.mark.anyio
async def test_browser_timestamp_cannot_backdate_access():
    from datetime import datetime, timezone
    from uuid import uuid4
    from unittest.mock import AsyncMock
    from app.services.analytics_service import AnalyticsService
    from app.schemas.analytics import AccessCreateRequest
    repository = AsyncMock()
    repository.get_or_create_visitor.return_value = SimpleNamespace(id=uuid4())
    repository.get_access.return_value = None
    payload = AccessCreateRequest(id=uuid4(), identity={'identity_kind': 'AUTHENTICATED', 'identifier': 'test'}, accessed_at=datetime(2000, 1, 1, tzinfo=timezone.utc))
    before = datetime.now(timezone.utc)
    await AnalyticsService(repository, identity_secret='test').record_access(payload)
    assert repository.create_access.call_args.args[2] >= before
