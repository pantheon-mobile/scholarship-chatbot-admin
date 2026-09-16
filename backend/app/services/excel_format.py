"""Download formatting from the 2026-09-16 customer form definition."""
import math
import unicodedata

from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils import get_column_letter

# Each column specifies (Excel width, wrap text, horizontal alignment).
def col(width, wrap=False, right=False):
    return (width, wrap, "right" if right else "left")

FAQ_COLUMNS = [col(9.3, right=True), col(59.3, True), col(79.3, True),
               *[col(39.3, True) for _ in range(10)],
               *[col(24.3, True) for _ in range(4)], col(14.3)]
FORMATS = {
    "データソース一覧": [col(9.3, right=True), col(9.3), col(59.3, True), col(59.3, True),
                    col(7.3), col(7.3), col(24.3, True), *[col(14.3, True) for _ in range(3)],
                    col(9.3, right=True), col(9.3, right=True), *[col(12.3) for _ in range(3)], col(19.3, right=True)],
    "URLリスト": [col(59.3), col(59.3, True)],
    "種別": [col(7.3), col(14.3, True), col(49.3, True)],
    "FAQ一覧": [*FAQ_COLUMNS, col(19.3, right=True)],
    "FAQ一括登録更新": FAQ_COLUMNS,
    "区分": [col(7.3, right=True), col(14.3, True), col(49.3, True)],
    "カテゴリ一覧": [col(9.3, right=True), *[col(49.3, True) for _ in range(3)]],
    "チャット履歴": [col(24.3, True), col(19.3, True), col(19.3, True), col(24.3, True),
                col(7.3, right=True), col(59.3, True), col(79.3, True), col(9.3), col(9.3),
                col(39.3, True), col(19.3, right=True), col(19.3, right=True)],
    "ユーザリスト": [col(19.3, True), col(19.3, True), col(29.3, True), col(19.3, right=True)],
}


def apply_download_format(sheet):
    columns = ([col(9.3, right=True), *[col(49.3, True) for _ in range(sheet.max_column - 1)]]
               if sheet.title == "カテゴリ一覧" else FORMATS[sheet.title])
    sheet.sheet_format.defaultRowHeight = 18
    sheet.row_dimensions[1].height = 18
    font = Font(name="游ゴシック", size=11, color="000000", bold=False)
    for index, (width, wrap, horizontal) in enumerate(columns, start=1):
        dimension = sheet.column_dimensions[get_column_letter(index)]
        dimension.width = width
        # Column defaults also format future input in the empty import templates.
        dimension.font = font
        dimension.alignment = Alignment(horizontal=horizontal, vertical="top", wrap_text=wrap)
        dimension.border = Border()
        for row in sheet.iter_rows(min_col=index, max_col=index, min_row=2):
            cell = row[0]
            cell.font = font
            cell.alignment = Alignment(horizontal=horizontal, vertical="top", wrap_text=wrap)
            cell.border = Border()
        header = sheet.cell(1, index)
        header.font = font
        header.fill = PatternFill(fill_type="solid", fgColor="F8F7F7")
        header.alignment = Alignment(horizontal="center", vertical="center")
        header.border = Border()

    for row in sheet.iter_rows(min_row=2):
        lines = 1
        for cell, (width, wrap, _) in zip(row, columns):
            if cell.value is None:
                continue
            # Estimate character widths for Japanese/Latin text, including explicit
            # line breaks. Excel's maximum row height is 409.5 points.
            text = str(cell.value).replace("\r\n", "\n").replace("\r", "\n")
            count = 0
            for line in text.split("\n"):
                units = sum(0 if unicodedata.combining(char) else 2 if unicodedata.east_asian_width(char) in "FWA" else 1 for char in line)
                count += max(1, math.ceil(units / max(1, width - 1))) if wrap else 1
            lines = max(lines, count)
        sheet.row_dimensions[row[0].row].height = min(409.5, 18 * lines)
