"""Excel work must leave the request event loop responsive."""
import asyncio
from io import BytesIO
from threading import Event, get_ident
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from fastapi import UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from app.schemas.faq import FaqFilters
from app.schemas.data_source import DataSourceFilters
from app.services import faq_service, data_source_service
from app.services.category_service import CategoryService
from app.services.classification_service import ClassificationService
from app.services.faq_classification_service import FaqClassificationService


async def assert_responsive(operation, monkeypatch, owner, name):
    started, release = Event(), Event()
    observations = []
    original = getattr(owner, name)
    loop_thread = get_ident()

    def slow(*args, **kwargs):
        started.set()
        observations.append((get_ident() != loop_thread, release.wait(2)))
        return original(*args, **kwargs)

    monkeypatch.setattr(owner, name, slow)
    task = asyncio.create_task(operation())
    try:
        async def wait_started():
            while not started.is_set():
                if task.done():
                    await task
                    pytest.fail('Excel operation was not reached')
                await asyncio.sleep(0.001)
        await asyncio.wait_for(wait_started(), timeout=5)
        release.set()
        result = await task
        assert observations and all(off_loop and resumed for off_loop, resumed in observations)
        return result
    finally:
        release.set()
        if not task.done():
            await task


def upload(headers):
    workbook = Workbook()
    workbook.active.append(headers)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return UploadFile(filename='test.xlsx', file=output)


@pytest.mark.anyio
@pytest.mark.parametrize('kind', ['faq', 'source'])
@pytest.mark.parametrize('phase', ['validate_excel_expansion', 'load_workbook', 'rows'])
async def test_import_checks_and_parses_without_blocking(kind, phase, monkeypatch):
    repo = AsyncMock()
    repo.list_import_classifications.return_value = []
    if kind == 'faq':
        module = faq_service
        service = module.FaqService(repo)
        headers = [*module.FAQ_IMPORT_FIXED_HEADERS, '区分1', '区分2', '区分3', '区分4', 'チャット利用']
        error = module.FaqError
    else:
        module = data_source_service
        service = module.DataSourceService(repo)
        headers = module.DATA_SOURCE_IMPORT_HEADERS
        error = module.DataSourceImportError
    source = upload(headers)
    async def operation():
        with pytest.raises(error) as caught:
            await service.import_excel(source)
        assert caught.value.code.endswith('_EMPTY')
        repo.commit.assert_not_awaited()
    owner, method = (ReadOnlyWorksheet, '_cells_by_row') if phase == 'rows' else (module, phase)
    await assert_responsive(operation, monkeypatch, owner, method)


@pytest.mark.anyio
@pytest.mark.parametrize('kind', ['faq', 'template', 'source', 'category', 'classification', 'faq_classification'])
async def test_exports_render_off_event_loop(kind, monkeypatch):
    repo = AsyncMock()
    repo.list.return_value = ([], 0, 0)
    repo.list_type_labels.return_value = {}
    repo.list_import_classifications.return_value = []
    repo.list_types.return_value = []
    repo.list_all.return_value = []
    if kind == 'faq':
        service = faq_service.FaqService(repo)
        service.validate_filters = AsyncMock(return_value={})
        operation = lambda: service.export_excel(FaqFilters(), {})
    elif kind == 'template':
        operation = faq_service.FaqService(repo).create_import_template
    elif kind == 'source':
        service = data_source_service.DataSourceService(repo)
        service.list = AsyncMock(return_value=SimpleNamespace(items=[], total_pages=1))
        operation = lambda: service.export_excel(DataSourceFilters())
    else:
        service = {'category': CategoryService, 'classification': ClassificationService,
                   'faq_classification': FaqClassificationService}[kind](repo)
        operation = service.export_excel
    content = await assert_responsive(operation, monkeypatch, Workbook, 'save')
    workbook = load_workbook(BytesIO(content))
    assert workbook.active.max_row == 1
    workbook.close()
