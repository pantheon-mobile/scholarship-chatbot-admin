from __future__ import annotations

from io import BytesIO
from pathlib import Path
from dataclasses import dataclass
from zipfile import BadZipFile, ZipFile, is_zipfile
from zoneinfo import ZoneInfo

from fastapi import UploadFile
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.models.data_source import DataSource
from app.models.category import Category
from app.repositories.data_source import DataSourceRepository
from app.schemas.data_source import (
    BulkDeleteRequest,
    ClassificationAssignment,
    DataSourceClassificationResponse,
    DataSourceCategoryResponse,
    DataSourceFileResponse,
    DataSourceFilters,
    DataSourceListResponse,
    DataSourceImportResponse,
    DataSourceImportRowError,
    DataSourceResponse,
    DataSourceWebsiteResponse,
    FileDataSourceUpdateRequest,
    WebsiteDataSourceCreateRequest,
    WebsiteDataSourceUpdateRequest,
)
from app.services.file_upload_validation import FileUploadValidationError, validate_uploads
from app.services.website_url_validation import WebsiteUrlValidationError, validate_website_url
from app.storage.base import StorageAdapter
from app.services.ingestion_processor import DataSourceCleanupProcessor


class DataSourceNotFoundError(Exception):
    pass


class DataSourceVersionConflictError(Exception):
    pass


class PageNotFoundError(Exception):
    pass


class ClassificationMismatchError(Exception):
    pass


class DataSourceCategoryNotFoundError(Exception):
    pass


class FileDataSourceRequiredError(Exception):
    pass


class DataSourceUpdateError(Exception):
    pass


class DataSourceCleanupError(Exception):
    pass


class WebsiteDataSourceCreateError(Exception):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


class WebsiteDataSourceRequiredError(Exception):
    pass


class WebsiteDataSourceUpdateError(Exception):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


class FileUploadError(Exception):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


DATA_SOURCE_IMPORT_HEADERS = ["ID", "種類", "タイトル", "ファイル名／URL", "形式", "状態", "カテゴリ", "種別1", "種別2", "種別3", "サイズ", "文字数", "回答ソース", "優先度", "参照元リンク", "更新日時"]
DATA_SOURCE_IMPORT_MAX_BYTES = 10 * 1024 * 1024
DATA_SOURCE_IMPORT_MAX_ROWS = 1000


class DataSourceImportError(Exception):
    def __init__(self, code: str, message: str, errors: list[DataSourceImportRowError] | None = None) -> None:
        super().__init__(message)
        self.code, self.message, self.errors = code, message, errors or []


class DataSourceService:
    def __init__(self, repository: DataSourceRepository, cleanup_processor: DataSourceCleanupProcessor | None = None) -> None:
        self.repository = repository
        self.cleanup_processor = cleanup_processor

    @staticmethod
    def category_paths(categories: list[Category]) -> dict[int, str]:
        by_id = {category.id: category for category in categories}
        paths: dict[int, str] = {}

        def resolve(category_id: int, visiting: set[int]) -> str:
            if category_id in paths:
                return paths[category_id]
            category = by_id[category_id]
            if category_id in visiting or category.parent_id is None or category.parent_id not in by_id:
                path = category.name
            else:
                path = f"{resolve(category.parent_id, visiting | {category_id})}/{category.name}"
            paths[category_id] = path
            return path

        for category_id in by_id:
            resolve(category_id, set())
        return paths

    @classmethod
    def serialize(cls, row: DataSource, categories: dict[int, Category], paths: dict[int, str]) -> DataSourceResponse:
        classifications = sorted(row.classification_links, key=lambda link: link.classification_type.display_order)
        row_category_id = getattr(row, "category_id", None)
        category = categories.get(row_category_id) if row_category_id is not None else None
        category_path = paths.get(category.id) if category else None
        return DataSourceResponse(
            id=row.id,
            source_type=row.source_type,
            title=row.title,
            format=row.format,
            status=row.status,
            category=DataSourceCategoryResponse(
                id=category.id,
                name=category.name,
                parent_id=category.parent_id,
                path=category_path or category.name,
            ) if category else None,
            category_name=category_path if category else row.category_name,
            size_bytes=row.size_bytes,
            character_count=row.character_count,
            answer_source_enabled=row.answer_source_enabled,
            priority=row.priority,
            reference_link_visible=row.reference_link_visible,
            updated_at=row.updated_at,
            version=row.version,
            file=DataSourceFileResponse(file_name=row.file.file_name) if row.file else None,
            website=DataSourceWebsiteResponse(
                url=row.website.url,
                last_fetched_at=getattr(row.website, "last_fetched_at", None),
            ) if row.website else None,
            classifications=[DataSourceClassificationResponse(
                type_code=link.classification_type.type_code,
                classification_type_id=link.classification_type_id,
                classification_value_id=link.classification_value_id,
                display_label=link.classification_type.display_label,
                value_name=link.classification_value.value_name,
            ) for link in classifications],
        )

    async def list(self, filters: DataSourceFilters) -> DataSourceListResponse:
        rows, total_count, total_pages, total_size = await self.repository.list(filters)
        if filters.page > 1 and filters.page > total_pages:
            raise PageNotFoundError()
        category_rows = await self.repository.list_categories()
        categories = {category.id: category for category in category_rows}
        paths = self.category_paths(category_rows)
        return DataSourceListResponse(
            items=[self.serialize(row, categories, paths) for row in rows],
            page=filters.page,
            page_size=filters.page_size,
            total_count=total_count,
            total_pages=total_pages,
            total_size_bytes=total_size,
            sort=filters.sort,
            order=filters.order,
        )

    async def _get(self, data_source_id: int) -> DataSource:
        row = await self.repository.get(data_source_id)
        if row is None:
            raise DataSourceNotFoundError()
        return row

    async def get(self, data_source_id: int) -> DataSourceResponse:
        row = await self._get(data_source_id)
        category_rows = await self.repository.list_categories()
        return self.serialize(row, {category.id: category for category in category_rows}, self.category_paths(category_rows))

    async def validate_category(self, category_id: int | None) -> None:
        if category_id is not None and not await self.repository.category_exists(category_id):
            raise DataSourceCategoryNotFoundError()

    async def update_file_attributes(
        self,
        data_source_id: int,
        payload: FileDataSourceUpdateRequest,
    ) -> DataSourceResponse:
        row = await self._get(data_source_id)
        if row.source_type != "FILE" or row.file is None:
            raise FileDataSourceRequiredError()
        await self.validate_category(payload.category_id)

        title = payload.title.strip() or row.file.file_name
        classifications: list[tuple[int, int]] = []
        for type_code, value_id in (
            ("TYPE_1", payload.type_1_value_id),
            ("TYPE_2", payload.type_2_value_id),
            ("TYPE_3", payload.type_3_value_id),
        ):
            if value_id is None:
                continue
            pair = await self.repository.resolve_classification_value(type_code, value_id)
            if pair is None:
                raise ClassificationMismatchError("種別値と種別の組み合わせが不正です。")
            classifications.append(pair)

        try:
            updated = await self.repository.update_file_attributes(
                data_source_id, payload, title, classifications
            )
        except Exception as exc:
            await self.repository.rollback()
            raise DataSourceUpdateError() from exc
        if not updated:
            raise DataSourceVersionConflictError()
        return await self.get(data_source_id)

    async def create_website_source(
        self,
        payload: WebsiteDataSourceCreateRequest,
    ) -> DataSourceResponse:
        try:
            url = validate_website_url(payload.url)
        except WebsiteUrlValidationError as exc:
            raise WebsiteDataSourceCreateError(exc.code, exc.message) from exc

        title = payload.title.strip() or url
        if len(title) > 500:
            raise WebsiteDataSourceCreateError("TITLE_TOO_LONG", "タイトルが長すぎます。")
        if payload.priority not in {"HIGH", "MEDIUM", "LOW"}:
            raise WebsiteDataSourceCreateError("INVALID_PRIORITY", "回答利用の優先度が不正です。")
        try:
            await self.validate_category(payload.category_id)
        except DataSourceCategoryNotFoundError as exc:
            raise WebsiteDataSourceCreateError("CATEGORY_NOT_FOUND", "指定されたカテゴリが存在しません。") from exc

        classifications: list[tuple[int, int]] = []
        for type_code, value_id in (
            ("TYPE_1", payload.type_1_value_id),
            ("TYPE_2", payload.type_2_value_id),
            ("TYPE_3", payload.type_3_value_id),
        ):
            if value_id is None:
                continue
            pair = await self.repository.resolve_classification_value(type_code, value_id)
            if pair is None:
                raise WebsiteDataSourceCreateError("INVALID_CLASSIFICATION", "種別値と種別の組み合わせが不正です。")
            classifications.append(pair)

        try:
            data_source_id = await self.repository.create_website_source(
                url=url,
                title=title,
                priority=payload.priority,
                answer_source_enabled=payload.answer_source_enabled,
                reference_link_visible=payload.reference_link_visible,
                category_id=payload.category_id,
                classifications=classifications,
            )
            await self.repository.commit()
        except Exception as exc:
            await self.repository.rollback()
            raise WebsiteDataSourceCreateError("WEB_DATA_SOURCE_CREATE_FAILED", "Webサイトの追加に失敗しました。") from exc
        return await self.get(data_source_id)

    async def update_website_attributes(
        self,
        data_source_id: int,
        payload: WebsiteDataSourceUpdateRequest,
    ) -> DataSourceResponse:
        row = await self._get(data_source_id)
        if row.source_type != "WEB" or row.website is None:
            raise WebsiteDataSourceRequiredError()
        try:
            await self.validate_category(payload.category_id)
        except DataSourceCategoryNotFoundError as exc:
            raise WebsiteDataSourceUpdateError("CATEGORY_NOT_FOUND", "指定されたカテゴリが存在しません。") from exc

        try:
            url = validate_website_url(payload.url)
        except WebsiteUrlValidationError as exc:
            raise WebsiteDataSourceUpdateError(exc.code, exc.message) from exc

        title = payload.title.strip() or url
        if len(title) > 500:
            raise WebsiteDataSourceUpdateError("TITLE_TOO_LONG", "タイトルが長すぎます。")
        if payload.priority not in {"HIGH", "MEDIUM", "LOW"}:
            raise WebsiteDataSourceUpdateError("INVALID_PRIORITY", "回答利用の優先度が不正です。")

        classifications: list[tuple[int, int]] = []
        for type_code, value_id in (
            ("TYPE_1", payload.type_1_value_id),
            ("TYPE_2", payload.type_2_value_id),
            ("TYPE_3", payload.type_3_value_id),
        ):
            if value_id is None:
                continue
            pair = await self.repository.resolve_classification_value(type_code, value_id)
            if pair is None:
                raise WebsiteDataSourceUpdateError("INVALID_CLASSIFICATION", "種別値と種別の組み合わせが不正です。")
            classifications.append(pair)

        try:
            updated = await self.repository.update_website_attributes(
                data_source_id, payload, url=url, title=title, classifications=classifications
            )
        except Exception as exc:
            await self.repository.rollback()
            raise WebsiteDataSourceUpdateError("WEB_DATA_SOURCE_UPDATE_FAILED", "Webサイトの更新に失敗しました。") from exc
        if not updated:
            raise DataSourceVersionConflictError()
        return await self.get(data_source_id)

    async def update_answer_source(self, data_source_id: int, enabled: bool, version: int) -> DataSourceResponse:
        await self._get(data_source_id)
        if not await self.repository.update_toggle(data_source_id, "answer_source_enabled", enabled, version):
            raise DataSourceVersionConflictError()
        return await self.get(data_source_id)

    async def update_reference_link(self, data_source_id: int, visible: bool, version: int) -> DataSourceResponse:
        await self._get(data_source_id)
        if not await self.repository.update_toggle(data_source_id, "reference_link_visible", visible, version):
            raise DataSourceVersionConflictError()
        return await self.get(data_source_id)

    async def delete(self, data_source_id: int, version: int) -> None:
        row = await self._get(data_source_id)
        if row.version != version:
            raise DataSourceVersionConflictError()
        if self.cleanup_processor:
            try:
                await self.cleanup_processor.cleanup([row])
            except Exception as exc:
                raise DataSourceCleanupError() from exc
        if not await self.repository.delete_one(data_source_id, version):
            raise DataSourceVersionConflictError()

    async def bulk_delete(self, payload: BulkDeleteRequest) -> int:
        if self.cleanup_processor:
            rows = [await self._get(target.id) for target in payload.items]
            versions = {row.id: row.version for row in rows}
            if any(versions.get(target.id) != target.version for target in payload.items):
                raise DataSourceVersionConflictError()
            try:
                await self.cleanup_processor.cleanup(rows)
            except Exception as exc:
                raise DataSourceCleanupError() from exc
        try:
            return await self.repository.bulk_delete(payload.items)
        except LookupError as exc:
            raise DataSourceNotFoundError() from exc
        except ValueError as exc:
            if str(exc) == "version_mismatch":
                raise DataSourceVersionConflictError() from exc
            raise

    async def validate_classification_assignments(self, assignments: list[ClassificationAssignment]) -> None:
        seen_types: set[int] = set()
        for assignment in assignments:
            if assignment.classification_type_id in seen_types:
                raise ClassificationMismatchError("同じ種別を複数指定できません。")
            seen_types.add(assignment.classification_type_id)
            if not await self.repository.classification_value_matches_type(
                assignment.classification_type_id,
                assignment.classification_value_id,
            ):
                raise ClassificationMismatchError("種別値と種別の組み合わせが不正です。")

    async def create_file_sources(
        self,
        files: list[UploadFile],
        storage: StorageAdapter,
        *,
        title: str | None,
        type_1_value_id: int | None,
        type_2_value_id: int | None,
        type_3_value_id: int | None,
        priority: str,
        answer_source_enabled: bool,
        reference_link_visible: bool,
        category_id: int | None = None,
    ) -> list[DataSourceResponse]:
        try:
            validated = validate_uploads(files)
        except FileUploadValidationError as exc:
            raise FileUploadError(exc.code, exc.message) from exc

        normalized_title = (title or "").strip()
        if len(validated) > 1 and normalized_title:
            raise FileUploadError("TITLE_NOT_ALLOWED", "複数ファイルを選択した場合はタイトルを指定できません。")
        if len(normalized_title) > 500:
            raise FileUploadError("INVALID_TITLE", "タイトルは500文字以内で入力してください。")
        if priority not in {"HIGH", "MEDIUM", "LOW"}:
            raise FileUploadError("INVALID_PRIORITY", "回答利用の優先度が不正です。")
        try:
            await self.validate_category(category_id)
        except DataSourceCategoryNotFoundError as exc:
            raise FileUploadError("CATEGORY_NOT_FOUND", "指定されたカテゴリが存在しません。") from exc

        classification_pairs: list[tuple[int, int]] = []
        for type_code, value_id in (
            ("TYPE_1", type_1_value_id),
            ("TYPE_2", type_2_value_id),
            ("TYPE_3", type_3_value_id),
        ):
            if value_id is None:
                continue
            pair = await self.repository.resolve_classification_value(type_code, value_id)
            if pair is None:
                raise FileUploadError("INVALID_CLASSIFICATION", "種別値と種別の組み合わせが不正です。")
            classification_pairs.append(pair)

        staged: list[tuple[Path, str]] = []
        finalized_keys: list[str] = []
        records: list[dict] = []
        try:
            for item in validated:
                storage_key = storage.create_storage_key(item.extension)
                temporary_path = storage.save_temporary(item.upload.file)
                staged.append((temporary_path, storage_key))
                records.append({
                    "title": normalized_title if len(validated) == 1 and normalized_title else item.file_name,
                    "file_name": item.file_name,
                    "storage_key": storage_key,
                    "extension": item.extension,
                    "size_bytes": item.size_bytes,
                    "content_type": item.content_type,
                })

            ids = await self.repository.create_file_sources(
                records,
                priority=priority,
                answer_source_enabled=answer_source_enabled,
                reference_link_visible=reference_link_visible,
                category_id=category_id,
                classifications=classification_pairs,
            )
            for temporary_path, storage_key in staged:
                storage.finalize(temporary_path, storage_key)
                finalized_keys.append(storage_key)
            await self.repository.commit()
        except Exception as exc:
            await self.repository.rollback()
            for temporary_path, _ in staged:
                storage.delete_temporary(temporary_path)
            for storage_key in finalized_keys:
                storage.delete(storage_key)
            if isinstance(exc, FileUploadError):
                raise
            raise FileUploadError("FILE_SAVE_FAILED", "ファイルの追加に失敗しました。") from exc

        result: list[DataSourceResponse] = []
        for data_source_id in ids:
            result.append(await self.get(data_source_id))
        return result

    async def export_excel(self, filters: DataSourceFilters) -> bytes:
        export_filters = filters.model_copy(update={"page": 1, "page_size": 100})
        all_rows: list[DataSourceResponse] = []
        while True:
            result = await self.list(export_filters)
            all_rows.extend(result.items)
            if export_filters.page >= result.total_pages:
                break
            export_filters = export_filters.model_copy(update={"page": export_filters.page + 1})

        source_labels = {"FILE": "ファイル", "WEB": "Web"}
        status_labels = {"PREPARING": "準備中", "TRAINING": "学習中", "AVAILABLE": "利用可", "ERROR": "エラー"}
        priority_labels = {"HIGH": "高", "MEDIUM": "中", "LOW": "低"}
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "データソース一覧"
        worksheet.append(DATA_SOURCE_IMPORT_HEADERS)
        for row in all_rows:
            values = {item.type_code: item.value_name for item in row.classifications}
            location = row.file.file_name if row.file else row.website.url if row.website else ""
            updated = row.updated_at.astimezone(ZoneInfo("Asia/Tokyo")).strftime("%Y/%m/%d %H:%M")
            worksheet.append([
                row.id, source_labels[row.source_type], row.title, location, row.format,
                status_labels[row.status], row.category_name or "", values.get("TYPE_1", ""),
                values.get("TYPE_2", ""), values.get("TYPE_3", ""), row.size_bytes,
                row.character_count, "有効" if row.answer_source_enabled else "無効",
                priority_labels[row.priority], "表示" if row.reference_link_visible else "非表示", updated,
            ])
        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    @staticmethod
    def _import_text(value: object) -> str:
        return "" if value is None else str(value).strip()

    @classmethod
    def _import_id(cls, value: object) -> int:
        if isinstance(value, bool):
            raise ValueError
        if isinstance(value, int) and value > 0:
            return value
        if isinstance(value, float) and value.is_integer() and value > 0:
            return int(value)
        if isinstance(value, str) and value.strip().isdigit() and int(value.strip()) > 0:
            return int(value.strip())
        raise ValueError

    async def import_excel(self, upload: UploadFile) -> DataSourceImportResponse:
        filename = upload.filename or ""
        if Path(filename).suffix.lower() != ".xlsx":
            raise DataSourceImportError("DATA_SOURCE_IMPORT_INVALID_FORMAT", "xlsx形式のファイルを選択してください。")
        content = await upload.read(DATA_SOURCE_IMPORT_MAX_BYTES + 1)
        if len(content) > DATA_SOURCE_IMPORT_MAX_BYTES:
            raise DataSourceImportError("DATA_SOURCE_IMPORT_FILE_TOO_LARGE", "ファイルサイズは10MB以下にしてください。")
        if not content or not is_zipfile(BytesIO(content)):
            raise DataSourceImportError("DATA_SOURCE_IMPORT_INVALID_FORMAT", "有効なxlsxファイルを選択してください。")
        try:
            with ZipFile(BytesIO(content)) as archive:
                names = {name.lower() for name in archive.namelist()}
                if any(name.endswith("vbaproject.bin") for name in names):
                    raise DataSourceImportError("DATA_SOURCE_IMPORT_INVALID_FORMAT", "マクロを含むExcelファイルは使用できません。")
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError):
            raise DataSourceImportError("DATA_SOURCE_IMPORT_INVALID_FORMAT", "有効なxlsxファイルを選択してください。") from None
        worksheet = workbook.active
        if worksheet.max_column != len(DATA_SOURCE_IMPORT_HEADERS):
            workbook.close()
            raise DataSourceImportError("DATA_SOURCE_IMPORT_INVALID_COLUMNS", "Excelの列数または列順が正しくありません。")
        header_cells = next(worksheet.iter_rows(min_row=1, max_row=1, max_col=len(DATA_SOURCE_IMPORT_HEADERS)))
        headers = [self._import_text(cell.value) for cell in header_cells]
        if headers != DATA_SOURCE_IMPORT_HEADERS or any(cell.data_type == "f" for cell in header_cells):
            workbook.close()
            raise DataSourceImportError("DATA_SOURCE_IMPORT_INVALID_COLUMNS", "Excelの列数または列順が正しくありません。")
        if worksheet.max_row - 1 > DATA_SOURCE_IMPORT_MAX_ROWS:
            workbook.close()
            raise DataSourceImportError("DATA_SOURCE_IMPORT_TOO_MANY_ROWS", "データ行は1000行以内にしてください。")

        raw_rows: list[tuple[int, list[object]]] = []
        errors: list[DataSourceImportRowError] = []
        ids: list[int] = []
        for row_number, cells in enumerate(worksheet.iter_rows(min_row=2, max_col=len(headers)), start=2):
            if all(self._import_text(cell.value) == "" for cell in cells):
                continue
            for index, cell in enumerate(cells):
                if cell.data_type == "f":
                    errors.append(DataSourceImportRowError(row=row_number, column=headers[index], code="FORMULA_NOT_ALLOWED", message="数式は入力できません。"))
            values = [cell.value for cell in cells]
            try:
                data_source_id = self._import_id(values[0])
                ids.append(data_source_id)
            except ValueError:
                errors.append(DataSourceImportRowError(row=row_number, column="ID", code="ID_INVALID", message="IDは必須の正の整数です。行の追加はできません。"))
            raw_rows.append((row_number, values))
        workbook.close()
        if not raw_rows:
            raise DataSourceImportError("DATA_SOURCE_IMPORT_EMPTY", "更新するデータがありません。")
        duplicates = {value for value in ids if ids.count(value) > 1}
        for row_number, values in raw_rows:
            try:
                if self._import_id(values[0]) in duplicates:
                    errors.append(DataSourceImportRowError(row=row_number, column="ID", code="ID_DUPLICATE", message="IDがExcel内で重複しています。"))
            except ValueError:
                pass

        existing_rows = await self.repository.get_for_update_many(list(set(ids)))
        existing = {int(row.id): row for row in existing_rows}
        categories = await self.repository.list_categories()
        category_paths = self.category_paths(categories)
        category_by_path = {path: category_id for category_id, path in category_paths.items()}
        type_definitions = await self.repository.list_import_classifications()
        types = {item.type_code: (int(item.id), {value.value_name: int(value.id) for value in item.values}) for item in type_definitions}
        source_labels = {"FILE": "ファイル", "WEB": "Web"}
        status_labels = {"PREPARING": "準備中", "TRAINING": "学習中", "AVAILABLE": "利用可", "ERROR": "エラー"}
        priority_by_label = {"高": "HIGH", "中": "MEDIUM", "低": "LOW"}
        updates: list[dict] = []
        immutable_indexes = (1, 3, 4, 5, 10, 11, 15)
        for row_number, values in raw_rows:
            try:
                data_source_id = self._import_id(values[0])
            except ValueError:
                continue
            row = existing.get(data_source_id)
            if row is None:
                errors.append(DataSourceImportRowError(row=row_number, column="ID", code="NOT_FOUND", message="指定されたデータソースが存在しません。行の追加はできません。"))
                continue
            current_classes = {link.classification_type.type_code: link.classification_value.value_name for link in row.classification_links}
            location = row.file.file_name if row.file else row.website.url if row.website else ""
            current = [row.id, source_labels[row.source_type], row.title, location, row.format, status_labels[row.status],
                       category_paths.get(row.category_id, ""), current_classes.get("TYPE_1", ""), current_classes.get("TYPE_2", ""), current_classes.get("TYPE_3", ""),
                       row.size_bytes, row.character_count, "有効" if row.answer_source_enabled else "無効",
                       {"HIGH": "高", "MEDIUM": "中", "LOW": "低"}[row.priority], "表示" if row.reference_link_visible else "非表示",
                       row.updated_at.astimezone(ZoneInfo("Asia/Tokyo")).strftime("%Y/%m/%d %H:%M")]
            for index in immutable_indexes:
                if self._import_text(values[index]) != self._import_text(current[index]):
                    errors.append(DataSourceImportRowError(row=row_number, column=headers[index], code="READ_ONLY_CHANGED", message="この項目は変更できません。"))
            title = self._import_text(values[2])
            if not title or len(title) > 500:
                errors.append(DataSourceImportRowError(row=row_number, column="タイトル", code="TITLE_INVALID", message="タイトルは1～500文字で入力してください。"))
            category_text = self._import_text(values[6])
            category_id = category_by_path.get(category_text) if category_text else None
            legacy_category_unchanged = row.category_id is None and category_text == self._import_text(row.category_name)
            if category_text and category_id is None and not legacy_category_unchanged:
                errors.append(DataSourceImportRowError(row=row_number, column="カテゴリ", code="CATEGORY_NOT_FOUND", message="指定されたカテゴリが存在しません。"))
            classifications = []
            for offset, code in enumerate(("TYPE_1", "TYPE_2", "TYPE_3"), start=7):
                label = self._import_text(values[offset])
                definition = types.get(code)
                value_id = definition[1].get(label) if definition and label else None
                if label and value_id is None:
                    errors.append(DataSourceImportRowError(row=row_number, column=headers[offset], code="CLASSIFICATION_NOT_FOUND", message="指定された種別が存在しません。"))
                elif value_id is not None:
                    classifications.append((definition[0], value_id))
            answer = self._import_text(values[12])
            priority = priority_by_label.get(self._import_text(values[13]))
            reference = self._import_text(values[14])
            if answer not in ("有効", "無効"):
                errors.append(DataSourceImportRowError(row=row_number, column="回答ソース", code="ANSWER_SOURCE_INVALID", message="「有効」または「無効」を入力してください。"))
            if priority is None:
                errors.append(DataSourceImportRowError(row=row_number, column="優先度", code="PRIORITY_INVALID", message="「高」「中」「低」のいずれかを入力してください。"))
            if reference not in ("表示", "非表示"):
                errors.append(DataSourceImportRowError(row=row_number, column="参照元リンク", code="REFERENCE_LINK_INVALID", message="「表示」または「非表示」を入力してください。"))
            proposed = dict(id=data_source_id, title=title, category_id=category_id, classifications=classifications,
                            answer_source_enabled=answer == "有効", priority=priority or row.priority,
                            reference_link_visible=reference == "表示")
            comparable = (title, category_id, sorted(classifications), answer == "有効", priority, reference == "表示")
            original = (row.title, row.category_id, sorted((int(link.classification_type_id), int(link.classification_value_id)) for link in row.classification_links), row.answer_source_enabled, row.priority, row.reference_link_visible)
            if comparable != original:
                updates.append(proposed)
        if errors:
            await self.repository.rollback()
            raise DataSourceImportError("DATA_SOURCE_IMPORT_VALIDATION_ERROR", "入力内容にエラーがあります。", errors)
        try:
            await self.repository.apply_import_updates(updates)
        except Exception as exc:
            await self.repository.rollback()
            raise DataSourceImportError("DATA_SOURCE_IMPORT_FAILED", "データソースの一括更新に失敗しました。") from exc
        return DataSourceImportResponse(updated_count=len(updates), processed_count=len(raw_rows))
