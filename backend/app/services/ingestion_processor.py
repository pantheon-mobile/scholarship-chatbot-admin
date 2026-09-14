from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
import csv
import io
import json
import logging
import os
from pathlib import PurePath
import time
from typing import Protocol

import boto3
import httpx
from docx import Document

from app.models.data_source import DataSource
from app.services.document_conversion import (
    convert_pdf,
    generate_pdf_content_metadata,
    convert_plain_text,
    convert_pptx,
    convert_xlsx,
    crawl_website,
)
from app.storage import LocalStorage, S3Storage


logger = logging.getLogger(__name__)


@dataclass(frozen=True)
class IngestionResult:
    character_count: int | None = None
    discovered_title: str | None = None


@dataclass(frozen=True)
class IngestionArtifact:
    name: str
    body: bytes
    content_type: str
    character_count: int | None = None
    source_url: str | None = None
    metadata: dict | None = None


class IngestionProcessor(Protocol):
    async def process(self, data_source: DataSource) -> IngestionResult: ...


class DataSourceCleanupProcessor(Protocol):
    async def cleanup(self, data_sources: list[DataSource]) -> None: ...


class HttpIngestionProcessor:
    """Adapter for the conversion/crawl/S3/Knowledge Base processing service."""

    def __init__(self, endpoint: str, *, timeout_seconds: float = 1800) -> None:
        self.endpoint = endpoint.rstrip("/")
        self.timeout_seconds = timeout_seconds

    async def process(self, data_source: DataSource) -> IngestionResult:
        payload = {
            "data_source_id": data_source.id,
            "source_type": data_source.source_type,
            "format": data_source.format,
            "title": data_source.title,
            "file": {
                "file_name": data_source.file.file_name,
                "storage_key": data_source.file.storage_key,
                "mime_type": data_source.file.mime_type,
            } if data_source.file else None,
            "website": {"url": data_source.website.url} if data_source.website else None,
            "metadata": AwsIngestionProcessor._data_source_metadata(data_source),
        }
        async with httpx.AsyncClient(timeout=self.timeout_seconds) as client:
            response = await client.post(f"{self.endpoint}/process", json=payload)
            response.raise_for_status()
            body = response.json()
        return IngestionResult(
            character_count=body.get("character_count"),
            discovered_title=body.get("discovered_title"),
        )


class LocalDataSourceCleanupProcessor:
    """Delete originals when running the self-contained local environment."""

    def __init__(self) -> None:
        self.storage = LocalStorage()

    async def cleanup(self, data_sources: list[DataSource]) -> None:
        for data_source in data_sources:
            if data_source.file and data_source.file.storage_key:
                self.storage.delete(data_source.file.storage_key)


class AwsIngestionProcessor:
    """Prepare a source, place Bedrock artifacts in S3, then synchronize its KB."""

    def __init__(self) -> None:
        self.region = os.getenv("AWS_REGION", "ap-northeast-1")
        self.bucket = os.getenv("INGESTION_S3_BUCKET", "").strip()
        if not self.bucket:
            raise RuntimeError("INGESTION_S3_BUCKET is required")
        self.s3 = boto3.client("s3", region_name=self.region)
        self.bedrock = boto3.client("bedrock-agent", region_name=self.region)
        if os.getenv("STORAGE_BACKEND", "local").lower() == "s3":
            self.source_storage = S3Storage(
                self.bucket,
                os.getenv("INGESTION_ORIGINAL_PREFIX", "documents/admin/originals/"),
            )
        else:
            self.source_storage = LocalStorage()

    async def process(self, data_source: DataSource) -> IngestionResult:
        kind = self._kind(data_source)
        knowledge_base_id, bedrock_data_source_id = self._kb_config(kind)
        crawl_report: dict | None = {} if kind == "WEB" else None
        artifacts = self._artifacts(data_source, kind, crawl_report=crawl_report)
        if not artifacts:
            raise RuntimeError("取り込み対象の文書が0件です。")
        discovered_title = None
        if crawl_report is not None and data_source.website is not None:
            root_url = data_source.website.url.rstrip("/")
            root_page = next((
                page for page in crawl_report.get("pages", [])
                if str(page.get("source_url", "")).rstrip("/") == root_url
            ), None)
            if root_page:
                discovered_title = str(root_page.get("title") or "").strip()[:500] or None
        effective_title = (
            discovered_title
            if discovered_title and data_source.website is not None and data_source.title == data_source.website.url
            else data_source.title
        )
        prefix = os.getenv(
            f"INGESTION_{kind}_S3_PREFIX",
            self._default_s3_prefix(kind),
        ).strip("/") + "/"
        source_prefix = f"{prefix}{data_source.id}/"
        self._clear_prefix(source_prefix)
        total_characters = 0
        has_character_count = False
        for artifact in artifacts:
            if artifact.character_count is not None:
                total_characters += artifact.character_count
                has_character_count = True
            key = f"{source_prefix}{artifact.name}"
            artifact_metadata = dict(artifact.metadata or {})
            management_metadata = self._data_source_metadata(data_source)
            if artifact_metadata.get("category") and management_metadata.get("category"):
                artifact_metadata["content_category"] = artifact_metadata.pop("category")
            attributes = {
                "data_source_id": str(data_source.id),
                "source_type": data_source.source_type,
                "source_format": data_source.format,
                "source_title": effective_title[:500],
                "answer_source_enabled": bool(data_source.answer_source_enabled),
                "answer_priority": data_source.priority,
                "reference_link_visible": bool(data_source.reference_link_visible),
                "ingestion_kind": kind,
                "source_url": artifact.source_url or "",
                "processed_at": datetime.now(timezone.utc).isoformat(),
                **artifact_metadata,
                **management_metadata,
                **self._artifact_identity_metadata(
                    data_source, artifact, effective_title, kind
                ),
            }
            attributes = {key: value for key, value in attributes.items() if value not in (None, "")}
            metadata = json.dumps(
                {"metadataAttributes": attributes}, ensure_ascii=False
            ).encode("utf-8")
            self.s3.put_object(
                Bucket=self.bucket, Key=key, Body=artifact.body,
                ContentType=artifact.content_type,
            )
            self.s3.put_object(
                Bucket=self.bucket, Key=f"{key}.metadata.json", Body=metadata,
                ContentType="application/json",
            )
        if crawl_report is not None:
            self._write_web_crawl_logs(data_source.id, crawl_report)
        self._synchronize(knowledge_base_id, bedrock_data_source_id)
        conversion_methods = sorted({
            str(artifact.metadata.get("conversion_method"))
            for artifact in artifacts if artifact.metadata.get("conversion_method")
        })
        selection_reasons = sorted({
            str(artifact.metadata.get("selection_reason"))
            for artifact in artifacts if artifact.metadata.get("selection_reason")
        })
        logger.info(
            "ingestion completed",
            extra={
                "data_source_id": data_source.id, "ingestion_kind": kind,
                "artifact_count": len(artifacts), "character_count": total_characters,
                "crawl_summary": crawl_report.get("summary") if crawl_report else None,
                "conversion_methods": conversion_methods,
                "selection_reasons": selection_reasons,
            },
        )
        return IngestionResult(
            character_count=total_characters if has_character_count else None,
            discovered_title=discovered_title,
        )

    @staticmethod
    def _data_source_metadata(data_source: DataSource) -> dict:
        """Build category/classification attributes shared by every artifact."""
        metadata: dict[str, str | int] = {}
        category_id = getattr(data_source, "category_id", None)
        if category_id is not None:
            metadata["category_id"] = int(category_id)
        category_path = str(
            getattr(data_source, "_ingestion_category_path", None)
            or getattr(data_source, "category_name", None)
            or ""
        ).strip()
        if category_path:
            metadata["category_path"] = category_path
            metadata["category"] = category_path

        links = getattr(data_source, "classification_links", None) or []
        for link in links:
            classification_type = getattr(link, "classification_type", None)
            classification_value = getattr(link, "classification_value", None)
            type_code = str(getattr(classification_type, "type_code", "")).upper()
            if type_code not in {"TYPE_1", "TYPE_2", "TYPE_3"}:
                continue
            suffix = type_code.removeprefix("TYPE_")
            value_id = getattr(link, "classification_value_id", None)
            value_name = str(getattr(classification_value, "value_name", "")).strip()
            if value_id is not None:
                metadata[f"type_{suffix}_id"] = int(value_id)
            if value_name:
                metadata[f"type_{suffix}"] = value_name
                metadata[f"type{suffix}"] = value_name
        metadata["answer_source"] = (
            "enabled" if bool(getattr(data_source, "answer_source_enabled", True)) else "disabled"
        )
        metadata["priority"] = str(getattr(data_source, "priority", "MEDIUM")).lower()
        return metadata

    @staticmethod
    def _artifact_identity_metadata(
        data_source: DataSource,
        artifact: IngestionArtifact,
        effective_title: str,
        kind: str,
    ) -> dict:
        original_name = (
            data_source.file.file_name
            if getattr(data_source, "file", None)
            else artifact.source_url or artifact.name
        )
        default_formats = {
            "PDF": "PDF_MARKDOWN",
            "WEB": "WEB_MARKDOWN",
            "EXCEL": "EXCEL_MARKDOWN",
            "WORD": "WORD_DOCX",
            "PPT": "PPT_MARKDOWN",
            "TEXT": "TEXT_MARKDOWN",
        }
        return {
            # Keep the proven PoC keys as aliases while retaining the management
            # application's canonical keys used by chat filtering.
            "datasource_id": str(data_source.id),
            "ingestion_format": (artifact.metadata or {}).get(
                "ingestion_format", default_formats[kind]
            ),
            "source_file_name": artifact.name,
            "original_source_file_name": PurePath(str(original_name)).name,
            "original_title": effective_title[:500],
            "generated_at": datetime.now(timezone.utc).isoformat(),
        }

    async def cleanup(self, data_sources: list[DataSource]) -> None:
        """Remove source-specific artifacts, refresh affected KBs, then originals."""
        sync_targets: set[tuple[str, str]] = set()
        original_keys: list[str] = []
        removal_prefixes: list[str] = []
        for data_source in data_sources:
            kind = self._kind(data_source)
            prefix = os.getenv(
                f"INGESTION_{kind}_S3_PREFIX", self._default_s3_prefix(kind)
            ).strip("/") + "/"
            removal_prefixes.append(f"{prefix}{data_source.id}/")
            sync_targets.add(self._kb_config(kind))
            if data_source.file and data_source.file.storage_key:
                original_keys.append(data_source.file.storage_key)

        # Validate every KB/DS setting before changing S3.
        for removal_prefix in removal_prefixes:
            self._clear_prefix(removal_prefix)

        # Bedrock removes vectors whose S3 source disappeared during this sync.
        for knowledge_base_id, data_source_id in sorted(sync_targets):
            self._synchronize(knowledge_base_id, data_source_id)

        # Keep originals until every affected KB has synchronized successfully so
        # an operator can retry/recover when AWS synchronization fails.
        for storage_key in original_keys:
            self.source_storage.delete(storage_key)

    def _clear_prefix(self, prefix: str) -> None:
        paginator = self.s3.get_paginator("list_objects_v2")
        for page in paginator.paginate(Bucket=self.bucket, Prefix=prefix):
            objects = [{"Key": item["Key"]} for item in page.get("Contents", [])]
            if objects:
                self.s3.delete_objects(Bucket=self.bucket, Delete={"Objects": objects})

    @staticmethod
    def _kind(data_source: DataSource) -> str:
        if data_source.source_type == "WEB":
            return "WEB"
        extension = data_source.format.lower().lstrip(".")
        kinds = {
            "pdf": "PDF", "xlsx": "EXCEL", "docx": "WORD", "pptx": "PPT",
            "txt": "TEXT", "csv": "TEXT",
        }
        if extension not in kinds:
            raise RuntimeError(f"夜間変換に未対応のファイル形式です: {extension}")
        return kinds[extension]

    @staticmethod
    def _default_s3_prefix(kind: str) -> str:
        # Bedrock Knowledge Bases currently allow at most five data sources per
        # knowledge base. Plain TXT/CSV artifacts therefore share the PDF data
        # source and its S3 prefix instead of consuming a sixth data source.
        prefix_kind = "PDF" if kind == "TEXT" else kind
        return f"documents/admin/kb-source/{prefix_kind.lower()}/"

    def _artifacts(
        self, data_source: DataSource, kind: str, *, crawl_report: dict | None = None
    ) -> list[IngestionArtifact]:
        if kind == "WEB":
            return self._markdown_artifacts(crawl_website(data_source.website.url, crawl_report))
        if data_source.file is None or not data_source.file.storage_key:
            raise RuntimeError("元ファイルの保存先がありません。")
        content = self.source_storage.read(data_source.file.storage_key)
        name = data_source.file.file_name
        extension = data_source.format.lower().lstrip(".")
        if extension == "docx":
            paragraph_count = table_count = headings = 0
            header_present = footer_present = False
            try:
                document = Document(io.BytesIO(content))
                paragraph_count = len(document.paragraphs)
                table_count = len(document.tables)
                headings = sum(
                    1 for paragraph in document.paragraphs
                    if paragraph.style and (paragraph.style.name or "").lower().startswith("heading")
                )
                header_present = any(
                    paragraph.text.strip()
                    for section in document.sections
                    for paragraph in section.header.paragraphs
                )
                footer_present = any(
                    paragraph.text.strip()
                    for section in document.sections
                    for paragraph in section.footer.paragraphs
                )
            except Exception:
                # Bedrock can ingest valid Office variants that python-docx may
                # not fully understand; structural metadata must not block them.
                logger.warning("Word structure metadata could not be extracted", exc_info=True)
            return [IngestionArtifact(
                name=PurePath(name).name,
                body=content,
                content_type=(
                    "application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.document"
                ),
                metadata={
                    "conversion_method": "original",
                    "ingestion_format": "WORD_DOCX",
                    "original_source_file_name": PurePath(name).name,
                    "source_file_type": "word",
                    "original_extension": ".docx",
                    "paragraph_count": paragraph_count,
                    "table_count": table_count,
                    "heading_count": headings,
                    "header_present": header_present,
                    "footer_present": footer_present,
                },
            )]
        if extension == "pdf":
            documents = convert_pdf(content, name)
            try:
                semantic_metadata = generate_pdf_content_metadata(content, name, max_pages=5)
            except Exception as exc:
                semantic_metadata = {
                    "metadata_generation_method": "claude_pdf_head",
                    "metadata_generation_status": "FAILED",
                    "metadata_generation_error": str(exc)[:500],
                }
            return self._markdown_artifacts(documents, common_metadata=semantic_metadata)
        if extension == "xlsx":
            documents = convert_xlsx(content, name)
        elif extension == "pptx":
            documents = convert_pptx(content, name)
        elif extension in {"txt", "csv"}:
            documents = convert_plain_text(content, name)
        elif extension != "pdf":
            raise AssertionError(f"Unsupported configured ingestion kind: {kind}")
        return self._markdown_artifacts(documents)

    def _write_web_crawl_logs(self, data_source_id: int, report: dict) -> None:
        """Store crawl audit files outside the Knowledge Base inclusion prefix."""
        log_prefix = os.getenv(
            "WEB_CRAWL_LOG_PREFIX", "documents/admin/crawl-logs/"
        ).strip("/") + f"/{data_source_id}/"
        manifest_key = f"{log_prefix}manifest.json"
        previous: dict[str, str] = {}
        try:
            body = self.s3.get_object(Bucket=self.bucket, Key=manifest_key)["Body"].read()
            previous_payload = json.loads(body)
            previous = {
                item["source_url"]: item.get("content_hash", "")
                for item in previous_payload.get("pages", [])
            }
        except Exception:
            # A missing or corrupt old manifest must not prevent a fresh crawl.
            previous = {}

        current = {item["source_url"]: item["content_hash"] for item in report.get("pages", [])}
        for item in report.get("pages", []):
            old_hash = previous.get(item["source_url"])
            item["change"] = "NEW" if old_hash is None else (
                "UNCHANGED" if old_hash == item["content_hash"] else "UPDATED"
            )
        report["deleted_candidates"] = sorted(set(previous) - set(current))
        report.setdefault("summary", {})["new"] = sum(item["change"] == "NEW" for item in report.get("pages", []))
        report["summary"]["updated"] = sum(item["change"] == "UPDATED" for item in report.get("pages", []))
        report["summary"]["unchanged"] = sum(item["change"] == "UNCHANGED" for item in report.get("pages", []))
        report["summary"]["deleted_candidates"] = len(report["deleted_candidates"])

        timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
        self.s3.put_object(
            Bucket=self.bucket, Key=manifest_key,
            Body=json.dumps({"pages": report.get("pages", [])}, ensure_ascii=False, indent=2).encode("utf-8"),
            ContentType="application/json",
        )
        self.s3.put_object(
            Bucket=self.bucket, Key=f"{log_prefix}{timestamp}/crawl-report.json",
            Body=json.dumps(report, ensure_ascii=False, indent=2).encode("utf-8"),
            ContentType="application/json",
        )
        for name in ("errors", "skipped"):
            output = io.StringIO()
            writer = csv.DictWriter(output, fieldnames=["url", "depth", "reason"])
            writer.writeheader()
            writer.writerows(report.get(name, []))
            self.s3.put_object(
                Bucket=self.bucket, Key=f"{log_prefix}{timestamp}/crawl-{name}.csv",
                Body=("\ufeff" + output.getvalue()).encode("utf-8"),
                ContentType="text/csv; charset=utf-8",
            )

    @staticmethod
    def _markdown_artifacts(documents, *, common_metadata: dict | None = None) -> list[IngestionArtifact]:
        return [IngestionArtifact(
            name=document.name,
            body=document.markdown.encode("utf-8"),
            content_type="text/markdown; charset=utf-8",
            character_count=len(document.markdown),
            source_url=document.source_url,
            metadata={**(document.metadata or {}), **(common_metadata or {})},
        ) for document in documents]

    @staticmethod
    def _kb_config(kind: str) -> tuple[str, str]:
        knowledge_base_id = os.getenv(f"INGESTION_{kind}_KNOWLEDGE_BASE_ID", "").strip()
        data_source_id = os.getenv(f"INGESTION_{kind}_DATA_SOURCE_ID", "").strip()
        if not knowledge_base_id or not data_source_id:
            raise RuntimeError(f"{kind}用のKnowledge Base IDまたはData Source IDが未設定です。")
        return knowledge_base_id, data_source_id

    def _synchronize(self, knowledge_base_id: str, data_source_id: str) -> None:
        response = self.bedrock.start_ingestion_job(
            knowledgeBaseId=knowledge_base_id,
            dataSourceId=data_source_id,
        )
        ingestion_job_id = response["ingestionJob"]["ingestionJobId"]
        timeout_seconds = int(os.getenv("KNOWLEDGE_BASE_SYNC_TIMEOUT_SECONDS", "1800"))
        deadline = time.monotonic() + timeout_seconds
        while time.monotonic() < deadline:
            current = self.bedrock.get_ingestion_job(
                knowledgeBaseId=knowledge_base_id,
                dataSourceId=data_source_id,
                ingestionJobId=ingestion_job_id,
            )["ingestionJob"]
            status = current["status"]
            if status == "COMPLETE":
                return
            if status in {"FAILED", "STOPPED"}:
                reasons = "; ".join(current.get("failureReasons", []))
                raise RuntimeError(f"Knowledge Base同期失敗: {reasons or status}")
            time.sleep(5)
        raise TimeoutError(f"Knowledge Base同期が{timeout_seconds}秒以内に完了しませんでした。")
